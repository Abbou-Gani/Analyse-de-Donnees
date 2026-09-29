import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COULEUR_ENTETE = "C15F3C"

LIBELLES = {
    "nombre_feuilles": "Nombre de feuilles",
    "nombre_lignes_total": "Nombre total de lignes",
    "nombre_lignes": "Nombre de lignes",
    "nombre_colonnes": "Nombre de colonnes",
    "colonnes_numeriques": "Colonnes chiffrées",
    "colonnes_texte": "Colonnes texte",
    "colonnes_dates": "Colonnes de date",
    "valeurs_manquantes": "Valeurs manquantes",
    "statistiques": "Statistiques",
    "nombre_paragraphes": "Nombre de paragraphes",
    "nombre_mots": "Nombre de mots",
    "nombre_tableaux": "Nombre de tableaux",
    "longueur_moyenne_paragraphe": "Longueur moyenne d'un paragraphe (mots)",
}

LIBELLES_STATS = {
    "count": "Valeurs",
    "mean": "Moyenne",
    "std": "Écart-type",
    "min": "Minimum",
    "25%": "25 %",
    "50%": "Médiane",
    "75%": "75 %",
    "max": "Maximum",
}


def _cellule(valeur):
    if valeur is None:
        return "—"
    if isinstance(valeur, float):
        return round(valeur, 4)
    return valeur


def _entete(ws, ligne: int, colonnes: list[str], largeurs: list[int]) -> None:
    for i, texte in enumerate(colonnes, start=1):
        cellule = ws.cell(row=ligne, column=i, value=texte)
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor=COULEUR_ENTETE)
        cellule.alignment = Alignment(vertical="center")
    for i, largeur in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.freeze_panes = ws.cell(row=ligne + 1, column=1)


def _feuille(wb, titre, colonnes, lignes, largeurs):
    ws = wb.create_sheet(titre)
    _entete(ws, 1, colonnes, largeurs)
    for r, valeurs in enumerate(lignes, start=2):
        for c, valeur in enumerate(valeurs, start=1):
            cellule = ws.cell(row=r, column=c, value=_cellule(valeur))
            if isinstance(valeur, str) and len(valeur) > 60:
                cellule.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def _lignes_resume(resume: dict) -> list[list]:
    lignes = []
    for cle, valeur in (resume or {}).items():
        if cle == "statistiques":
            continue
        libelle = LIBELLES.get(cle, cle.replace("_", " ").capitalize())
        if isinstance(valeur, dict):
            for sous_cle, sous_valeur in valeur.items():
                lignes.append([libelle, str(sous_cle), _cellule(sous_valeur)])
        elif isinstance(valeur, list):
            lignes.append([libelle, "", ", ".join(str(v) for v in valeur) or "—"])
        else:
            lignes.append([libelle, "", _cellule(valeur)])
    return lignes


def _lignes_stats(resume: dict) -> list[list]:
    stats = (resume or {}).get("statistiques") or {}
    ordre = ["count", "mean", "std", "min", "25%", "50%", "75%", "max"]
    lignes = []
    for colonne, valeurs in stats.items():
        lignes.append(
            [colonne] + [_cellule(valeurs.get(cle)) for cle in ordre]
        )
    return lignes


def _lignes_anomalies(anomalies: list) -> list[list]:
    lignes = []
    for a in anomalies:
        exemples = ", ".join(str(v) for v in (a.get("valeurs") or [])[:5])
        lignes.append(
            [
                a.get("feuille") or "",
                a.get("colonne"),
                a.get("nombre_valeurs_aberrantes"),
                a.get("pourcentage"),
                a.get("borne_basse"),
                a.get("borne_haute"),
                a.get("moyenne"),
                a.get("mediane"),
                exemples,
            ]
        )
    return lignes


def _lignes_drill_down(anomalies: list) -> list[list]:
    lignes = []
    for a in anomalies:
        colonne = a.get("colonne")
        feuille = a.get("feuille") or ""
        for ligne in a.get("lignes") or []:
            numero = ligne.get("ligne")
            valeur = ligne.get(colonne)
            contexte = " · ".join(
                f"{cle} : {v}"
                for cle, v in ligne.items()
                if cle not in ("ligne", colonne) and v is not None
            )
            lignes.append([feuille, colonne, numero, valeur, contexte])
    return lignes


def _lignes_tendances(tendances: list) -> list[list]:
    return [
        [
            t.get("feuille") or "",
            t.get("colonne"),
            t.get("colonne_date"),
            t.get("sens"),
            t.get("valeur_premiere"),
            t.get("valeur_derniere"),
            t.get("evolution_pourcentage"),
            t.get("valeur_min"),
            t.get("valeur_max"),
            t.get("valeur_moyenne"),
            t.get("nombre_points"),
        ]
        for t in tendances
    ]


def _lignes_recommandations(recommandations: list) -> list[list]:
    return [
        [
            r.get("priorite") or "",
            r.get("titre") or "",
            r.get("detail") or "",
            r.get("action") or "",
        ]
        for r in recommandations
    ]


def construire_export(
    nom_fichier: str,
    resume: dict,
    anomalies: list,
    tendances: list,
    recommandations: list,
    insights: str | None = None,
) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    ws = _feuille(
        wb,
        "Résumé",
        ["Indicateur", "Détail", "Valeur"],
        _lignes_resume(resume),
        [38, 30, 60],
    )
    ligne = ws.max_row + 2
    ws.cell(row=ligne, column=1, value="Fichier analysé").font = Font(bold=True)
    ws.cell(row=ligne, column=3, value=nom_fichier)
    if insights:
        ligne += 2
        ws.cell(row=ligne, column=1, value="Synthèse").font = Font(bold=True)
        cellule = ws.cell(row=ligne + 1, column=1, value=insights)
        cellule.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(
            start_row=ligne + 1, start_column=1, end_row=ligne + 6, end_column=3
        )

    stats = _lignes_stats(resume)
    if stats:
        _feuille(
            wb,
            "Statistiques",
            ["Colonne"] + list(LIBELLES_STATS.values()),
            stats,
            [28] + [14] * len(LIBELLES_STATS),
        )

    if anomalies:
        _feuille(
            wb,
            "Anomalies",
            [
                "Feuille",
                "Colonne",
                "Valeurs aberrantes",
                "% de la colonne",
                "Borne basse",
                "Borne haute",
                "Moyenne",
                "Médiane",
                "Exemples de valeurs",
            ],
            _lignes_anomalies(anomalies),
            [18, 24, 18, 16, 14, 14, 14, 14, 40],
        )
        drill = _lignes_drill_down(anomalies)
        if drill:
            _feuille(
                wb,
                "Lignes concernées",
                ["Feuille", "Colonne", "Ligne du fichier", "Valeur", "Contexte de la ligne"],
                drill,
                [18, 24, 16, 16, 70],
            )

    if tendances:
        _feuille(
            wb,
            "Tendances",
            [
                "Feuille",
                "Colonne",
                "Colonne de date",
                "Sens",
                "1re valeur",
                "Dernière valeur",
                "Évolution %",
                "Minimum",
                "Maximum",
                "Moyenne",
                "Points",
            ],
            _lignes_tendances(tendances),
            [18, 22, 18, 12, 14, 16, 14, 12, 12, 14, 10],
        )

    if recommandations:
        _feuille(
            wb,
            "Recommandations",
            ["Priorité", "Recommandation", "Détail", "Action proposée"],
            _lignes_recommandations(recommandations),
            [12, 42, 70, 60],
        )

    sortie = io.BytesIO()
    wb.save(sortie)
    return sortie.getvalue()
