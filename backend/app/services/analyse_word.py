import io
import re
from collections import Counter

from docx import Document


def analyser_word(contenu: bytes) -> dict:
    document = Document(io.BytesIO(contenu))

    paragraphes = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    texte_complet = " ".join(paragraphes)

    tableaux = []
    for tableau in document.tables:
        lignes = []
        for ligne in tableau.rows:
            lignes.append([cellule.text.strip() for cellule in ligne.cells])
        tableaux.append(lignes)

    mots_cles = _extraire_mots_cles(texte_complet)
    entites = _extraire_entites(texte_complet)
    resume = _generer_resume(paragraphes)
    nombre_mots = len(texte_complet.split())
    longueur_moyenne = (
        round(sum(len(p.split()) for p in paragraphes) / len(paragraphes), 1)
        if paragraphes
        else 0
    )
    recommandations = _generer_recommandations(
        paragraphes, tableaux, mots_cles, entites, nombre_mots, longueur_moyenne
    )

    return {
        "type": "word",
        "resume": {
            "nombre_paragraphes": len(paragraphes),
            "nombre_mots": nombre_mots,
            "nombre_tableaux": len(tableaux),
            "longueur_moyenne_paragraphe": longueur_moyenne,
        },
        "resume_texte": resume,
        "mots_cles": mots_cles,
        "entites": entites,
        "tableaux": tableaux,
        "anomalies": [],
        "tendances": [],
        "recommandations": recommandations,
    }


def _extraire_mots_cles(texte: str, nombre: int = 10) -> list[dict]:
    mots = re.findall(r"[a-zA-ZàâäéèêëîïôöùûüçÀÂÄÉÈÊËÎÏÔÖÙÛÜÇ']{3,}", texte.lower())
    mots_vides = {
        "les", "des", "une", "que", "qui", "quoi", "pour", "avec", "dans", "sur",
        "pas", "plus", "tout", "tous", "toute", "mais", "comme", "entre", "ces",
        "leur", "leurs", "vous", "nous", "ils", "elles", "cette", "cela", "fait",
        "etre", "avoir", "son", "ses", "est", "sont", "aux", "par", "chez",
        "the", "and", "for", "with", "this", "that", "from", "are", "was",
    }
    compteurs = Counter(m for m in mots if m not in mots_vides)
    return [{"mot": mot, "occurrences": n} for mot, n in compteurs.most_common(nombre)]


def _extraire_entites(texte: str) -> dict:
    montants = re.findall(r"\d+[.,]?\d*\s*€|\d+[.,]?\d*\s*euros?", texte, re.IGNORECASE)
    dates = re.findall(r"\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b|\b\d{4}-\d{2}-\d{2}\b", texte)
    nombres = re.findall(r"\b\d+(?:[.,]\d+)?\b", texte)
    return {
        "montants": montants,
        "dates": dates,
        "nombres_importants": nombres[:20],
    }


def _generer_resume(paragraphes: list[str]) -> str:
    if not paragraphes:
        return "Document vide."
    premieres_phrases = []
    for paragraphe in paragraphes[:3]:
        phrases = re.split(r"(?<=[.!?])\s+", paragraphe)
        premieres_phrases.append(phrases[0])
    base = " ".join(premieres_phrases)
    nb_mots = sum(len(p.split()) for p in paragraphes)
    return (
        f"Document de {len(paragraphes)} paragraphe(s) et {nb_mots} mot(s). "
        f"Ouverture : {base}"
    )


def _generer_recommandations(
    paragraphes: list[str],
    tableaux: list,
    mots_cles: list,
    entites: dict,
    nombre_mots: int,
    longueur_moyenne: float,
) -> list[dict]:
    recommandations = []
    nombre_paragraphes = len(paragraphes)

    if not paragraphes:
        recommandations.append(
            {
                "titre": "Document sans texte exploitable",
                "detail": (
                    "Aucun paragraphe non vide n'a été lu dans le fichier Word. "
                    "Le contenu est peut-être uniquement dans des images ou des objets."
                ),
                "action": (
                    "Collez le texte directement dans le document ou exportez-le "
                    "en version éditable, puis relancez l'analyse."
                ),
                "priorite": "haute",
            }
        )
    elif nombre_mots < 50:
        recommandations.append(
            {
                "titre": "Document très court",
                "detail": (
                    f"{nombre_paragraphes} paragraphe(s) et seulement {nombre_mots} mot(s). "
                    "L'analyse thématique a peu de matière."
                ),
                "action": (
                    "Complétez le contenu ou joignez les annexes avant de tirer "
                    "des conclusions sur les thèmes."
                ),
                "priorite": "moyenne",
            }
        )
    else:
        recommandations.append(
            {
                "titre": "Lisibilité et structure du document",
                "detail": (
                    f"{nombre_paragraphes} paragraphe(s), {nombre_mots} mot(s), "
                    f"en moyenne {longueur_moyenne} mot(s) par paragraphe. "
                    + (
                        "Des paragraphes très longs peuvent nuire à la lecture."
                        if longueur_moyenne > 80
                        else "La longueur moyenne des paragraphes est correcte."
                    )
                ),
                "action": (
                    "Gardez un titre par idée, des paragraphes courts et une "
                    "conclusion actionnable en fin de document."
                ),
                "priorite": "info",
            }
        )

    if tableaux:
        tailles = ", ".join(
            f"tableau {i + 1} : {len(t)} ligne(s)" for i, t in enumerate(tableaux[:4])
        )
        if len(tableaux) > 4:
            tailles += f"… ({len(tableaux)} au total)"
        recommandations.append(
            {
                "titre": "Exploiter les tableaux du document",
                "detail": (
                    f"{len(tableaux)} tableau(x) détecté(s) ({tailles}). "
                    "Ces données peuvent être croisées avec les chiffres Excel "
                    "si vous les exportez (copier → Excel)."
                ),
                "action": (
                    "Exportez les tableaux dans un classeur pour calculer moyennes, "
                    "totaux et écarts automatiquement."
                ),
                "priorite": "moyenne",
            }
        )
    else:
        recommandations.append(
            {
                "titre": "Ajouter une vue chiffrée",
                "detail": (
                    "Aucun tableau détecté dans le document : les constats restent "
                    "qualitatifs pour l'instant."
                ),
                "action": (
                    "Ajoutez un tableau de synthèse (indicateurs clés, avant/après) "
                    "pour appuyer l'analyse."
                ),
                "priorite": "info",
            }
        )

    if mots_cles:
        top = ", ".join(f"« {m['mot']} » ({m['occurrences']}×)" for m in mots_cles[:5])
        recommandations.append(
            {
                "titre": "Axes thématiques à approfondir",
                "detail": (
                    f"Mots les plus fréquents hors termes courants : {top}. "
                    f"Ils résument l'objet du document sur {nombre_mots} mot(s) analysés."
                ),
                "action": (
                    "Vérifiez que chaque thème a une recommandation ou une décision "
                    "associée dans votre compte rendu."
                ),
                "priorite": "info",
            }
        )

    montants = entites.get("montants") or []
    dates = entites.get("dates") or []
    if montants:
        exemples = ", ".join(montants[:5])
        recommandations.append(
            {
                "titre": "Contrôler les montants cités",
                "detail": (
                    f"{len(montants)} mention(s) d'argent détectée(s), par ex. {exemples}. "
                    "Ces chiffres doivent rester alignés avec le fichier Excel de référence."
                ),
                "action": (
                    "Recoupez chaque montant avec le classeur et signalez les écarts "
                    "avant diffusion."
                ),
                "priorite": "haute",
            }
        )
    if dates:
        exemples = ", ".join(dates[:5])
        recommandations.append(
            {
                "titre": "Vérifier la cohérence des dates",
                "detail": (
                    f"{len(dates)} date(s) repérée(s), par ex. {exemples}. "
                    "Des dates incohérentes faussent le récit chronologique."
                ),
                "action": (
                    "Uniformisez le format des dates et confirmez qu'elles "
                    "correspondent à la période analysée."
                ),
                "priorite": "moyenne",
            }
        )

    if not recommandations:
        recommandations.append(
            {
                "titre": "Document lisible",
                "detail": "Analyse textuelle terminée sans point d'attention particulier.",
                "action": "Conservez ce document comme référence pour vos prochains rapports.",
                "priorite": "info",
            }
        )
    return recommandations
